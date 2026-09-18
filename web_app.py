"""
Beacon Regression DB - Web App
================================
Run: streamlit run web_app.py
Opens at: http://localhost:8501
"""

import streamlit as st
import pandas as pd
import random
import io
import zipfile
from datetime import datetime, timedelta
from openpyxl import Workbook

# ============================================================
# PAGE CONFIG
# ============================================================
st.set_page_config(
    page_title="Beacon Regression Generator",
    page_icon="📊",
    layout="wide"
)

# ============================================================
# CONSTANTS & DEFAULTS
# ============================================================
INST_RULES = {
    "Term Loan":                     {"rev": False, "sub": True,  "curr": "INR"},
    "Working Capital Demand Loan":   {"rev": False, "sub": True,  "curr": "INR"},
    "Inter Company Deposit":         {"rev": False, "sub": False, "curr": "INR"},
    "Direct Assignment":             {"rev": False, "sub": False, "curr": "INR"},
    "External Commercial Borrowing": {"rev": False, "sub": True,  "curr": "USD"},
    "Re-Finance":                    {"rev": False, "sub": True,  "curr": "INR"},
    "Line Of Credit":                {"rev": True,  "sub": True,  "curr": "INR"},
    "Short Term Loan":               {"rev": True,  "sub": False, "curr": "INR"},
    "Cash Credit":                   {"rev": True,  "sub": True,  "curr": "INR"},
    "Overdraft":                     {"rev": True,  "sub": False, "curr": "INR"},
}

DEFAULT_LENDERS = [
    {"is_bank":"Yes","bank_id":"BNK-001","short":"HDFC","full":"HDFC Bank Ltd","loc":"Mumbai","pan":"AAACH1234E","grp":"HDFC Group","cls":"Bank","typ":"Private Bank","sub":"Scheduled Commercial Bank"},
    {"is_bank":"Yes","bank_id":"BNK-002","short":"AXIS","full":"Axis Bank Ltd","loc":"Mumbai","pan":"AAACA1234H","grp":"Axis Group","cls":"Bank","typ":"Private Bank","sub":"Scheduled Commercial Bank"},
    {"is_bank":"Yes","bank_id":"BNK-003","short":"SBI","full":"State Bank of India","loc":"Mumbai","pan":"AAACS1234G","grp":"SBI Group","cls":"Bank","typ":"PSU Bank","sub":"Scheduled Commercial Bank"},
    {"is_bank":"Yes","bank_id":"BNK-004","short":"KOTAK","full":"Kotak Mahindra Bank","loc":"Mumbai","pan":"AAACK1234I","grp":"Kotak Group","cls":"Bank","typ":"Private Bank","sub":"Scheduled Commercial Bank"},
    {"is_bank":"No","bank_id":"","short":"INFY","full":"Infosys Ltd","loc":"Bangalore","pan":"AAACI1234B","grp":"Infosys Group","cls":"Corporate","typ":"NBFC","sub":"NBFC-ICC"},
    {"is_bank":"No","bank_id":"","short":"TCS","full":"Tata Consultancy Services Ltd","loc":"Mumbai","pan":"AAACT1234C","grp":"Tata Group","cls":"Corporate","typ":"NBFC","sub":"NBFC-ICC"},
    {"is_bank":"No","bank_id":"","short":"RIL","full":"Reliance Industries Ltd","loc":"Mumbai","pan":"AAACR1234D","grp":"Reliance Group","cls":"Corporate","typ":"NBFC","sub":"NBFC-ICC"},
    {"is_bank":"No","bank_id":"","short":"XYZF","full":"XYZ Finance Ltd","loc":"Bangalore","pan":"AAAXYZ1234K","grp":"XYZ Group","cls":"NBFC","typ":"NBFC","sub":"NBFC-ICC"},
]

DCC = ["Actual/365","Actual/360","Actual/Actual By Payment Date","US 30/360","European 30/360"]
COUPON_TYPES = ["Fixed","Floating"]
COUPON_FREQ = ["1M","3M","6M","12M"]
TDS_TYPES = ["No TDS","Normal Tax Rate","Tax Rate"]
FEE_TYPES = ["Processing Fee","Advisory / Arranger Fee","Renewal Fee","Stamp Duty","Bank Charges","Commitment Fee"]
INDICES = ["Repo Rate","MCLR","LIBOR","SOFR","T-Bill"]

def add_days(d, n):
    if isinstance(d, str): d = datetime.strptime(d, "%Y-%m-%d")
    return (d + timedelta(days=n)).strftime("%Y-%m-%d")

def pick(l): return random.choice(l)

# ============================================================
# GENERATOR FUNCTIONS
# ============================================================
def gen_lender_wb(lenders):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Basic Client Details")
    ws.append(["Is Bank","Bank ID","Short Name","Full Name","Location","PAN","Group",
               "Related Party","Related Party Dropdown","Lender Class","Lender Type",
               "Lender Sub Type","TAN","Legal Entity Identifier","ID"])
    for i, l in enumerate(lenders, 1):
        ws.append([
            l.get("is_bank", "Yes"),
            l.get("bank_id", ""),
            l.get("short", ""),
            l.get("full", ""),
            l.get("loc", ""),
            l.get("pan", ""),
            l.get("grp", ""),
            "No","Parent - Holding",
            l.get("cls", "Bank"),
            l.get("typ", "Private Bank"),
            l.get("sub", "Scheduled Commercial Bank"),
            f"TAN{i:05d}",f"LEI{i:05d}",f"LEN-{i:03d}"
        ])
    ws2 = wb.create_sheet("Demat Details")
    ws2.append(["DP Name","DP ID","Demat Account","Scheme Name","Client ID","Depository","ID"])
    for i, l in enumerate(lenders, 1):
        short_name = l.get("short") or "Lender"
        ws2.append([f"{short_name} DP",f"DP{i:04d}",f"DEMAT{i:06d}","Scheme A",f"CL{i:05d}","NSDL",f"LEN-{i:03d}"])
    ws3 = wb.create_sheet("Contact Details")
    ws3.append(["Contact Person","Contact Number","Landline Number","Email Address","Designation","ID"])
    for i, l in enumerate(lenders, 1):
        ws3.append([f"Contact {i}",f"98{i:08d}",f"080-2{i:07d}",f"contact{i}@test.com","Manager",f"LEN-{i:03d}"])
    ws4 = wb.create_sheet("Bank Account Details")
    ws4.append(["PAN","Account Name","Name","Branch","Address","IFSC Code","MICR Code","Account Use","Account Number","ID"])
    for i, l in enumerate(lenders, 1):
        short_name = l.get("short") or "LEND"
        full_name = l.get("full", "")
        loc = l.get("loc") or "Main"
        pan = l.get("pan", "")
        ws4.append([pan, full_name, full_name, f"{loc} Main", f"{loc} Address",
                    f"{short_name[:4].upper():4s}000{i:04d}", f"56024{i:05d}", "Current", f"{10000+i}", f"LEN-{i:03d}"])
    ws5 = wb.create_sheet("Address Segment Details")
    ws5.append(["Office Location Type","Mark as Default","Office DUNS Number","Address Line 1",
                "Address Line 2","Address Line 3","City / Town","District",
                "State / Union Territory","Country","Pin Code","GSTIN","Mobile Number",
                "Email Address","Telephone Area Code","Telephone Number(s)",
                "Fax Area Code","Fax Number(s)","ID"])
    for i, l in enumerate(lenders, 1):
        loc = l.get("loc") or "Mumbai"
        pan = l.get("pan") or "AAACH1234E"
        ws5.append(["Registered Office","Yes",f"DUNS{i:05d}",loc,"Line 2","",loc,loc,
                    "Maharashtra" if loc=="Mumbai" else "Karnataka","India",
                    "400001" if loc=="Mumbai" else "560001",
                    f"{pan[:5]}1234E1Z5" if len(pan)>=5 else "27XXXXX1Z5",
                    f"98{i:08d}",f"info{i}@test.com","080",f"2{i:07d}","080",f"2{i:07d}",f"LEN-{i:03d}"])
    return wb

def gen_sanction_wb(n_sanctions, fy, start, company_name, lenders, instruments, inst_rules=None):
    if inst_rules is None:
        inst_rules = INST_RULES
    wb = Workbook()
    wb.remove(wb.active)
    basic, main, sub = [], [], []
    for i in range(n_sanctions):
        l = lenders[i % len(lenders)]
        sref = f"SM-{fy}-{str(i+1).zfill(5)}"
        picked = random.sample(instruments, min(random.choice([1,2,2,3]), len(instruments)))
        basic.append({"ref":sref,"lender":l.get("full", ""),"date":add_days(start,i*2),"company":company_name})
        for j, inst in enumerate(picked, 1):
            rule = inst_rules.get(inst, {"rev": False, "sub": True, "curr": "INR"})
            nature = pick(["Revolving","Non-Revolving"]) if rule["rev"] else "Non-Revolving"
            amt = pick([1000000,2500000,5000000,7500000,10000000])
            has_sub = "Yes" if (rule["sub"] and random.random()<0.4) else "No"
            main.append({"ref":sref,"no":j,"inst":inst,"nat":nature,"amt":amt,
                         "valid":add_days(start,360),"has_sub":has_sub,"curr":rule["curr"]})
            if has_sub=="Yes":
                sub.append({"ref":sref,"no":j,"inst":inst,"nat":nature,"amt":amt,
                            "valid":add_days(start,360),"curr":rule["curr"]})

    ws = wb.create_sheet("Basic Details")
    ws.append(["Sanction Reference*","Lender Full Name*","Sanction Date*","Company Full Name"])
    for b in basic: ws.append([b["ref"],b["lender"],b["date"],b["company"]])

    ws2 = wb.create_sheet("Main Breakup Details")
    ws2.append(["Sanction Reference*","Main Breakup Number(Used for mapping sub limits)",
                "Instrument*","Nature of breakup*","Sanction Amount*",
                "Availability/Validity Date*","Has Sub Limits*","Currency*"])
    for m in main:
        ws2.append([m["ref"],m["no"],m["inst"],m["nat"],m["amt"],m["valid"],m["has_sub"],m["curr"]])

    ws3 = wb.create_sheet("Sub Breakup Details")
    ws3.append(["Sanction Reference*","Main Breakup Number(Used for mapping sub limits)",
                "Instrument*","Nature of breakup*","Sanction Amount",
                "Availability/Validity Date*","Currency*"])
    for s in sub:
        ws3.append([s["ref"],s["no"],s["inst"],s["nat"],s["amt"],s["valid"],s["curr"]])

    ws4 = wb.create_sheet("dropdowns")
    ws4.append(["Nature of Sanction","Has Sub Limits","Instruments","","Currency"])
    for row in [["Non-Revolving","Yes","Term Loan","","USD"],
                ["Revolving","No","Working Capital Demand Loan","","JPY"],
                ["","","Inter Company Deposit","","EUR"],
                ["","","Direct Assignment","","INR"],
                ["","","External Commercial Borrowing","","GBP"],
                ["","","Re-Finance","","CHF"],
                ["","","Line Of Credit","","QAR"],
                ["","","Short Term Loan","","CAD"],
                ["","","Cash Credit","","SEK"],
                ["","","Overdraft","","SGD"]]:
        ws4.append(row)
    return wb

def gen_wcdl_wb(n_deals, fy, company_name, lenders):
    wb = Workbook()
    wb.remove(wb.active)
    sanctions = [{"ref":f"SM-{fy}-{str(i+1).zfill(5)}","lender":lenders[i%len(lenders)],
                  "facility":f"WC-Rev-{str(i+1).zfill(4)}"} for i in range(n_deals)]
    deals = []
    for i in range(n_deals):
        s = sanctions[i]
        ctype = pick(COUPON_TYPES)
        settle = add_days("2026-02-01", i*3)
        lender_full = s["lender"].get("full", "")
        deals.append({
            "ID":f"WC-{str(i+1).zfill(3)}","Is_Rollover":pick(["Yes","No"]),
            "Rollover_Type":pick(["With Interest","Without Interest"]),
            "Settlement_Date":settle,"Lender":lender_full,
            "Curr_Acc":str(10001+(i*2)),"Loan_Acc":str(20001+(i*2)),
            "Loan_Acc_No":f"LAN{i+1:08d}","Nature":pick(["Secured","Unsecured"]),
            "Security_Cover":pick([0,1,1,1]),"Company":company_name,
            "Sanction_ID":s["ref"],"Facility_ID":s["facility"],
            "Disb_Date":settle,"Disb_Amt":pick([20000,50000,100000,150000]),
            "Repay_Date":add_days(settle,90),"Coupon_Type":ctype,
            "Coupon_Freq":pick(COUPON_FREQ),
            "Fixed_Rate":round(random.uniform(0.06,0.12),4) if ctype=="Fixed" else "",
            "Index":pick(INDICES) if ctype=="Floating" else "",
            "Spread":round(random.uniform(0.01,0.05),4) if ctype=="Floating" else "",
            "DCC":pick(DCC),"First_Int":add_days(settle,3),
            "Next_Review":add_days(settle,33),"TDS":pick(TDS_TYPES),
        })

    ws = wb.create_sheet("Facility Details")
    ws.append(["ID*","Is Rollover","Rollover Type","Rollover ID","Settlement Date*",
               "Lender Full Name*","Current Bank Account*","Loan Bank Account",
               "Loan Account Number","Nature","Details Of Security","Security Cover",
               "Company Full Name*","Unsecured Loan Description",
               "Scheme Name (Re-finance)","Scheme Details (Re-finance)"])
    for d in deals:
        ws.append([d["ID"],d["Is_Rollover"],
                   d["Rollover_Type"] if d["Is_Rollover"]=="Yes" else "","",
                   d["Settlement_Date"],d["Lender"],d["Curr_Acc"],d["Loan_Acc"],
                   d["Loan_Acc_No"],d["Nature"],"",d["Security_Cover"],d["Company"],
                   "Unsecured Loan" if d["Nature"]=="Unsecured" else "","None",""])

    for sheet_name, headers, rows in [
        ("Ratings", ["ID*","Rating Reference*","Remarks"],
         [[d["ID"], pick(["AAA","AA+","AA","A+","Unrated"]), "Sample"] for d in deals]),
        ("Contact Master", ["ID*","Contact Person Name","Comments"],
         [[d["ID"], f"Contact {d['ID']}", "Sample"] for d in deals]),
        ("Sanction Details", ["ID*","Sanction ID*","Sanction Facility ID*","Is EMI*",
                              "EMI Override DCC","Maturity for EMI","PRINCIPAL MORATORIUM INSTALLMENT"],
         [[d["ID"],d["Sanction_ID"],d["Facility_ID"],"No","","",""] for d in deals]),
        ("Disbursement Schedule", ["ID*","Date*","Amount*","Payment Instructions","Remarks"],
         [[d["ID"],d["Disb_Date"],d["Disb_Amt"],"No","Auto"] for d in deals]),
        ("Repayment Schedule", ["ID*","Date*","Payment Date","Forward Rate (ECB)",
                                "Forward Date(ECB)","Amount*","Invoice Number","Remarks"],
         [[d["ID"],d["Repay_Date"],d["Repay_Date"],"","",d["Disb_Amt"],f"INV-{d['ID']}","Auto"] for d in deals]),
        ("Interest details", ["ID*","Period (Coupon Start Date)*","Coupon Type*",
                              "Coupon Payment Frequency*","First Interest Due Date*",
                              "Fixed Rate","Next Rate Review Date","Index (Benchmark Name)",
                              "Spread","DCC*"],
         [[d["ID"],d["Settlement_Date"],d["Coupon_Type"],d["Coupon_Freq"],d["First_Int"],
           d["Fixed_Rate"],d["Next_Review"],d["Index"],d["Spread"],d["DCC"]] for d in deals]),
        ("TDS", ["ID*","TDS","Start Date"],
         [[d["ID"],d["TDS"],d["Settlement_Date"]] for d in deals]),
        ("Pre Payment Penalty", ["ID*","Upto Months","No Prepayment","Applicable Rate(%)","Remarks"],
         [[d["ID"],pick([3,6,9,12]),pick(["Yes","No"]),pick([1.0,1.5,2.0]),"Auto"] for d in deals]),
    ]:
        wsx = wb.create_sheet(sheet_name)
        wsx.append(headers)
        for r in rows: wsx.append(r)

    ws10 = wb.create_sheet("Fees,Charges")
    ws10.append(["ID*","Invoice Date","Payment Due Date","Type","Amount","% of Disbursement",
                 "Counterparty Type","Counterparty Name","Is GST Applicable","GSTIN",
                 "Include for IndAS","TDS Type","Is RCM Applicable","Invoice Num",
                 "Details","Bank Account","Remarks"])
    for d in deals:
        for j in range(random.choice([1,2])):
            ws10.append([d["ID"],d["Disb_Date"],add_days(d["Disb_Date"],30),
                         pick(FEE_TYPES),pick([200,500,1000,2000,5000]),
                         round(random.uniform(0.05,1.0),2),
                         pick(["Lender","Arranger","Other"]),d["Lender"],
                         pick(["Yes","No"]),"27AAACH1234E1Z5",pick(["Yes","No"]),
                         pick(TDS_TYPES),pick(["Yes","No"]),f"INV-{d['ID']}-{j+1}",
                         "Auto",d["Curr_Acc"],""])
    return wb

def wb_to_bytes(wb):
    """Convert openpyxl Workbook to bytes for download"""
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()

# ============================================================
# STREAMLIT UI
# ============================================================
st.title("📊 Beacon Regression DB Generator")
st.markdown("Fill the form below to generate regression test files for Beacon Lending System.")

st.sidebar.header("⚙️ Configuration")

# --- Basic Settings ---
st.sidebar.subheader("Basic Settings")
company_name = st.sidebar.text_input("Company Full Name", value="Wipro Ltd")
n_sanctions = st.sidebar.number_input("Number of Sanctions", min_value=1, max_value=100, value=12)
n_wcdl = st.sidebar.number_input("Number of WCDL Deals", min_value=0, max_value=50, value=6)
fiscal_year = st.sidebar.number_input("Fiscal Year (e.g. 26 for 2026)", min_value=20, max_value=50, value=26)
start_date = st.sidebar.text_input("Start Date (YYYY-MM-DD)", value="2026-01-05")
random_seed = st.sidebar.number_input("Random Seed", min_value=1, max_value=99999, value=42)

# Prepare session state for editable Lenders and Instruments DataFrames
if "lenders_df" not in st.session_state:
    st.session_state["lenders_df"] = pd.DataFrame(DEFAULT_LENDERS)

if "inst_df" not in st.session_state:
    inst_list = []
    for inst, r in INST_RULES.items():
        inst_list.append({
            "Instrument": inst,
            "Revolving": r["rev"],
            "Has Sub Limits": r["sub"],
            "Currency": r["curr"]
        })
    st.session_state["inst_df"] = pd.DataFrame(inst_list)

# --- Lenders Dropdown Selection in Sidebar ---
st.sidebar.subheader("Lenders Selection")
all_lender_names = st.session_state["lenders_df"]["full"].tolist()
selected_lender_names = st.sidebar.multiselect(
    "Select Lenders:",
    options=all_lender_names,
    default=all_lender_names
)

# --- Instruments Dropdown Selection in Sidebar ---
st.sidebar.subheader("Instruments Selection")
all_inst_names = st.session_state["inst_df"]["Instrument"].tolist()
selected_inst_names = st.sidebar.multiselect(
    "Select Instruments:",
    options=all_inst_names,
    default=all_inst_names
)

# --- Main Area ---
col1, col2, col3 = st.columns(3)
with col1:
    st.metric("Lenders Selected", len(selected_lender_names))
with col2:
    st.metric("Instruments Selected", len(selected_inst_names))
with col3:
    st.metric("Sanctions to Generate", n_sanctions)

st.markdown("---")

# --- Interactive Editable Tables Section ---
st.subheader("✏️ Editable Configuration & Data")
st.markdown("You can directly edit lender details or instrument rules below, or add new ones:")

tab1, tab2 = st.tabs(["🏢 Editable Lenders Master", "📜 Editable Instruments Master"])

with tab1:
    st.markdown("Edit existing lenders or add new rows directly in the table:")
    edited_lenders_df = st.data_editor(
        st.session_state["lenders_df"],
        num_rows="dynamic",
        column_config={
            "is_bank": st.column_config.SelectboxColumn("Is Bank", options=["Yes", "No"], default="Yes"),
            "bank_id": st.column_config.TextColumn("Bank ID"),
            "short": st.column_config.TextColumn("Short Name"),
            "full": st.column_config.TextColumn("Full Name"),
            "loc": st.column_config.TextColumn("Location"),
            "pan": st.column_config.TextColumn("PAN"),
            "grp": st.column_config.TextColumn("Group"),
            "cls": st.column_config.SelectboxColumn("Lender Class", options=["Bank", "Corporate", "NBFC"], default="Bank"),
            "typ": st.column_config.TextColumn("Lender Type"),
            "sub": st.column_config.TextColumn("Lender Sub Type"),
        },
        key="lender_editor"
    )
    st.session_state["lenders_df"] = edited_lenders_df

with tab2:
    st.markdown("Edit instrument properties (Revolving, Sub Limits, Currency) or add custom instruments:")
    edited_inst_df = st.data_editor(
        st.session_state["inst_df"],
        num_rows="dynamic",
        column_config={
            "Instrument": st.column_config.TextColumn("Instrument Name"),
            "Revolving": st.column_config.CheckboxColumn("Revolving"),
            "Has Sub Limits": st.column_config.CheckboxColumn("Has Sub Limits"),
            "Currency": st.column_config.SelectboxColumn("Currency", options=["INR", "USD", "EUR", "GBP", "JPY", "CHF", "CAD", "SGD", "SEK", "QAR"], default="INR"),
        },
        key="inst_editor"
    )
    st.session_state["inst_df"] = edited_inst_df

st.markdown("---")

st.subheader("📋 Configuration Summary")
summary = pd.DataFrame([
    {"Setting": "Company", "Value": str(company_name)},
    {"Setting": "Sanctions", "Value": str(n_sanctions)},
    {"Setting": "WCDL Deals", "Value": str(n_wcdl)},
    {"Setting": "Fiscal Year", "Value": str(fiscal_year)},
    {"Setting": "Start Date", "Value": str(start_date)},
    {"Setting": "Lenders", "Value": f"{len(selected_lender_names)} selected"},
    {"Setting": "Instruments", "Value": f"{len(selected_inst_names)} selected"},
])
st.table(summary)

st.markdown("---")

# Filter selected lenders and instruments dictionary
active_lenders = [
    row.to_dict() for _, row in st.session_state["lenders_df"].iterrows()
    if row["full"] in selected_lender_names
]

active_inst_rules = {}
for _, row in st.session_state["inst_df"].iterrows():
    if row["Instrument"] in selected_inst_names:
        active_inst_rules[str(row["Instrument"])] = {
            "rev": bool(row.get("Revolving", False)),
            "sub": bool(row.get("Has Sub Limits", False)),
            "curr": str(row.get("Currency", "INR"))
        }

active_inst_names = list(active_inst_rules.keys())

# --- Generate Button ---
if st.button("🚀 Generate Files", type="primary", use_container_width=True):
    if not active_lenders or not active_inst_names:
        st.error("⚠️ Please select at least 1 lender and 1 instrument.")
    else:
        with st.spinner("Generating files... Please wait..."):
            random.seed(random_seed)

            # Generate all 3 workbooks using editable & selected parameters
            lender_wb = gen_lender_wb(active_lenders)
            sanction_wb = gen_sanction_wb(n_sanctions, fiscal_year, start_date,
                                          company_name, active_lenders, active_inst_names,
                                          inst_rules=active_inst_rules)
            wcdl_wb = gen_wcdl_wb(n_wcdl, fiscal_year, company_name, active_lenders)

            # Create ZIP
            zip_buf = io.BytesIO()
            with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                zf.writestr("Lender_Regression_Filled.xlsx", wb_to_bytes(lender_wb))
                zf.writestr("Sanction_Regression_Filled.xlsx", wb_to_bytes(sanction_wb))
                zf.writestr("WCDL_Regression_Filled.xlsx", wb_to_bytes(wcdl_wb))
            zip_buf.seek(0)

        st.success("✅ Files generated successfully!")
        st.balloons()

        st.markdown("### 📥 Download Files")

        c1, c2, c3 = st.columns(3)
        with c1:
            st.download_button(
                "⬇️ Lender File",
                data=wb_to_bytes(lender_wb),
                file_name="Lender_Regression_Filled.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        with c2:
            st.download_button(
                "⬇️ Sanction File",
                data=wb_to_bytes(sanction_wb),
                file_name="Sanction_Regression_Filled.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
        with c3:
            st.download_button(
                "⬇️ WCDL File",
                data=wb_to_bytes(wcdl_wb),
                file_name="WCDL_Regression_Filled.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

        st.markdown("### 📦 Download All (ZIP)")
        st.download_button(
            "⬇️ Download All Files (ZIP)",
            data=zip_buf.getvalue(),
            file_name="Beacon_Regression_Files.zip",
            mime="application/zip",
            use_container_width=True
        )

st.markdown("---")
st.caption("Beacon Regression Generator v1.0 | Made for Quantum Phinance")
